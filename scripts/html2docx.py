#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""HTML 报告 → Word 结构化转换模块。

把报告页的标题层级、正文段落、列表、表格提取为可编辑的 Word 元素，
图表位置按 manifest 插入已导出的高清 PNG。

版面约定（2026-10-09 修订）：
  1. 表格：首行若为「标题行」（只有首格有字、其余空或被合并），该行渲染为
     表标题（整行合并、居中、加粗、浅底），次行才是真表头（加粗白色 + 深底 +
     跨页重复）；`<tbody>` 下的裸 `<th>`（不在 `<tr>` 内）归并成表头行，
     避免整行表头丢失。
  2. 卡片：网页「标题 + 数字」的卡片组一律渲染成 Word 表格（见 CARD_* 系列），
     不再拉平成一行一行的清单。

用法（被 assemble_docx.py 调用）：
  from html2docx import new_document, convert_report, load_figmap
"""
import os
import re

from bs4 import BeautifulSoup
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

H1, H2, H3, H4 = 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'
CN_FONT = '宋体'

HEAD_FILL = '1A3A5F'      # 表头底色
TITLE_FILL = 'DCE6F2'     # 表标题底色
ZEBRA_FILL = 'F4F6F9'     # 隔行底色
CARD_TITLE_FILL = 'EAEFF6'

# 需要整体丢弃的容器（导航/页脚/脚本/交互控件）
SKIP_CLASS = {
    'nav-bar', 'nav-inner', 'nav-logo', 'nav-links', 'nav-dropdown', 'dropdown-menu',
    'dropdown-item', 'dropdown-divider', 'badge', 'footer', 'drill-hint', 'toc-item',
    'tab-btn', 'rank-tab-btn', 'tabs', 'street-tabs', 'no-print', 'modal', 'drawer',
    'legend-note', 'chart-controls', 'print-hide',
    # 网页专有装饰：导航、完成进度条、图例、交互提示
    'sub-nav-bar', 'sub-nav-inner', 'sub-nav-back', 'sub-nav-links', 'sub-nav-tab',
    'progress-wrap', 'progress-title', 'progress-bar', 'progress-fill',
    'trend-legend', 'legend-item', 'legend-dot', 'bar-track',
    'cat-icon', 'model-arrow', 'flow-arrow',
    # 浮动目录面板 / 年度切换标签 / 本地文件链接（不得进入 Word）
    'toc-panel', 'toc-float', 'toc-header', 'toc-body', 'toc-sub-wrap', 'toc-group',
    'toc-sub', 'toc-close', 'toc-toggle', 'toc-mask',
    'year-tabs', 'year-tab', 'excel-link', 'card-link', 'attach-title',
    'nav', 'legend', 'legend-wrap', 'foot', 'flow-arrow-line',
}
SKIP_TAG = {'script', 'style', 'noscript', 'button', 'svg', 'canvas', 'iframe', 'nav'}

# 行内元素：在 walk 中不单独成段，需与其相邻行内兄弟合并为一段
INLINE_TAG = {'span', 'a', 'strong', 'b', 'em', 'i', 'u', 'sub', 'sup', 'small',
              'code', 'mark', 'br', 'abbr', 'cite', 'time', 'label'}

# 原子卡片：内部虽为块级结构，但应合并成一段，避免被拆成零碎单行
ATOMIC_CLASS = {
    'quality-item', 'mini-card', 'summary-card', 'metric',
    'insight-card', 'note-card', 'highlight-card', 'cc-body', 'ac-body',
}
BLOCK_TAG = {'div', 'section', 'article', 'main', 'ul', 'ol', 'li', 'p', 'table', 'tr',
             'td', 'th', 'thead', 'tbody', 'figure', 'blockquote', 'header', 'footer',
             'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'dl', 'dt', 'dd', 'pre'}

# ---------------------------------------------------------------- 卡片规格
# 卡片字段类名 → 表头文字（None 表示不进入表格）
FIELD_LABEL = {
    # 名称类
    'label': '项目', 'k': '项目', 'lab': '项目', 'kpi-label': '项目',
    'stat-label': '项目', 'mc-month': '月份', 'y': '年度', 'ext-t': '项目',
    'src-name': '数据源', 'fc-name': '维度', 'criteria-label': '项目',
    # 数值类
    'num': '数值', 'v': '数值', 'val': '数值', 'kpi-val': '数值', 'stat-value': '数值',
    'mc-total': '数值', 'ext-v': '内容', 'src-meta': '规模与层级', 'fc-val': '构成',
    'criteria-value': '内容', 'mi-desc': '规模',
    # 备注类
    'tag': '备注', 'kpi-sub': '备注', 'stat-sub': '构成', 'stat-change': '变化',
    'mc-changes': '环比 / 同比', 'mc-comm': '涉及小区', 'avg': '月均', 'change': '同比',
    'match-rate': '关联状态', 'fc-desc': '说明', 'action': '处置', 'fc-icon': None,
    'mi-name': '数据源', 'mi-icon': None, 'src-icon': None, 'cat-icon': None,
    'pct': '占比',
    # ---- 高风险评分模型：维度卡
    'dim-name': '评分维度', 'dim-weight': '权重', 'dim-desc': '评分规则',
    # ---- 根因模型：因子行
    'factor-name': '影响因素', 'factor-bar-fill': '相关系数', 'factor-weight': '权重',
    # ---- 月报明细：指标项
    'ds-label': '指标', 'ds-val': '数值',
    # ---- 去重方法论：策略卡 / 流程图节点
    's-name': '识别策略', 's-detail': '识别逻辑',
    # ---- 底部抬升：矛盾分类卡 / 投诉类别条 / 月度柱组
    'contra-name': '矛盾分类', 'contra-num': '小区数', 'contra-desc': '涵盖内容',
    'contra-change': '2026年1-8月同比',
    'cat-name': '投诉类别', 'cat-val': '件数',
    'bar-label': '区间', 'bar-val': '每千户投诉量',
    # ---- 进度条（占比构成）
    'progress-bar-label': '构成',
    # ---- 风险等级 / 其他
    'mc-change': None, 'topic-num': None, 'topic-title': None,
}
# 列顺序（未列出的排最后）
FIELD_ORDER = ['label', 'k', 'lab', 'kpi-label', 'stat-label', 'criteria-label',
               'mc-month', 'y', 'ext-t', 'src-name', 'fc-name', 'mi-name',
               'ds-label', 'bar-label', 'cat-name', 'contra-name', 'factor-name',
               'dim-name', 's-name',
               'num', 'v', 'val', 'kpi-val', 'stat-value', 'mc-total',
               'src-meta', 'ext-v', 'fc-val', 'criteria-value', 'mi-desc',
               'ds-val', 'bar-val', 'cat-val', 'contra-num', 'factor-bar-fill',
               'dim-weight', 'factor-weight',
               'tag', 'kpi-sub', 'stat-sub', 'stat-change', 'mc-changes', 'mc-comm',
               'avg', 'change', 'match-rate', 'fc-desc', 'action',
               'pct', 'dim-desc', 'contra-desc', 's-detail', 'contra-change',
               'progress-bar-label']

# 扁平指标卡（同类连续出现 → 合并成一张表）
METRIC_CARD_CLASS = {
    'kpi', 'kpi-card', 'stat-card', 'stat-item', 'month-card', 'year-card',
    'ext-card', 'source-card', 'factor-card', 'model-item', 'criteria-box',
    # 新增：评分维度卡 / 根因因子行 / 月报指标项 / 策略卡 / 矛盾卡 / 类别条 / 柱组
    'model-dim', 'factor-row', 'ds-item', 'flow-strategy', 'contra-card',
    'cat-item', 'bar-row', 'progress-bar-wrap',
}
# 内容卡：类名 → (标题选择器, 正文选择器列表, 两列表头)
CONTENT_CARD_CLASS = {
    'finding-card': ('finding-title', ['finding-desc'], ('分析发现', '说明')),
    'issue-card': ('issue-title', ['issue-desc', 'issue-solution'], ('数据问题', '说明与解决方案')),
    'advice-card': ('ac-title', ['ac-desc'], ('治理建议', '内容')),
    'action-card': ('ac-title', ['ac-desc'], ('治理举措', '内容')),
    'case-card': ('cc-title', ['cc-desc'], ('典型案例', '说明')),
    'root-cause': ('rc-title', ['rc-desc'], ('根本原因', '说明')),
    'insight-box': ('ib-title', ['ib-content'], ('维度分析', '结论')),
    'field-cat': ('fc-title', [], ('字段类别', '字段清单')),
}
# 单一卡片（标题 + 正文）类名 → 选择器
SINGLE_CARD_CLASS = {
    'summary-section': ('summary-title', ['summary-text']),
    'suggestion-block': ('sug-header', ['sug-desc']),
}
# 提示框（浅底 + 左侧色条）→ 保留为带底色的单格表
CALLOUT_CLASS = {'highlight', 'callout', 'note', 'note-box', 'insight'}
# 小标题类名 → 提升为卡片标题
CARD_TITLE_CLASS = {
    'mini-table-title', 'summary-title', 'detail-title', 'model-title',
    'chart-title', 'map-title', 'sug-header', 'table-title',
}
# 脚注/图注类名 → 小字说明段
CAPTION_CLASS = {'table-note', 'wc-note', 'map-caption', 'chart-note', 'note-text',
                 'footnote', 'source-note', 'sub'}
# 预格式化文本（保留换行与空格）
MONO_CLASS = {'code-block', 'pre-block', 'ascii-table'}
# 关键词云 → 归并成一段关键词罗列
KEYWORD_CLASS = {'keyword-cloud'}
# 图表容器（CSS 柱状图，无导出图片时改渲染为数据表）
BARCHART_CLASS = {'bar-chart'}

# 严重度 emoji → 文字颜色（源报告用彩色圆点表达等级，Word 里改用颜色）
SEV_COLOR = {
    '\U0001F534': (0xC6, 0x28, 0x28),   # 红
    '\U0001F7E0': (0xE6, 0x51, 0x00),   # 橙
    '\U0001F7E1': (0xB2, 0x6A, 0x00),   # 黄
    '\U0001F7E2': (0x2E, 0x7D, 0x32),   # 绿
    '\U0001F535': (0x15, 0x65, 0xC0),   # 蓝
}

CALLOUT_FILL = 'EEF4FB'   # 提示框底色
MONO_FILL = 'F5F5F5'      # 预格式化块底色

# 正文可用宽度（cm）：A4 21 - 左右页边距 3.17*2
TEXT_W = 14.6


# ---------------------------------------------------------------- 样式
def new_document():
    from docx import Document
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(3.17)
        s.right_margin = Cm(3.17)
    st = doc.styles['Normal']
    st.font.name = CN_FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    for name, size, bold in [(H1, 18, True), (H2, 15, True), (H3, 13, True), (H4, 12, True)]:
        try:
            ss = doc.styles[name]
            ss.font.size = Pt(size)
            ss.font.bold = bold
            ss.font.name = CN_FONT
            ss.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5F)
            ss.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
        except Exception:
            pass
    return doc


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), color)
    tcPr.append(sh)


def vmerge_none(cell):
    """取消单元格的纵向合并（避免沿用模板残留）。"""
    tcPr = cell._tc.get_or_add_tcPr()
    vm = tcPr.find(qn('w:vMerge'))
    if vm is not None:
        tcPr.remove(vm)


def _run(cell, text, bold=False, size=9, color=None, align='left'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(str(text))
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = CN_FONT
    r.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    if color:
        r.font.color.rgb = RGBColor(*color)


def set_cell(cell, text, bold=False, color=None, align='left', size=9):
    _run(cell, text, bold, size, color, align)


def repeat_header(row):
    """把某行标记为跨页重复的表头行。"""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def add_body(doc, text, indent=True, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = CN_FONT
    r.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    return p


def add_card_title(doc, text, size=11):
    """卡片内的小标题：加粗、不缩进、与上文留白。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = CN_FONT
    r.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5F)
    r.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = CN_FONT
    r.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r.font.name = CN_FONT
    r.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    return p


def add_image(doc, path, width_cm=TEXT_W, caption=None):
    if not os.path.exists(path):
        add_body(doc, '［图表文件缺失：%s］' % os.path.basename(path), indent=False)
        return None
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        add_caption(doc, caption)


# ---------------------------------------------------------------- 表格
def _table_rows(tbl):
    """按文档顺序取表格行；`<tbody>` 下的裸 `<th>`（不在 tr 内）归并成头行。

    返回 [{'cells': [...], 'head': bool}]。
    """
    rows = []
    for el in tbl.descendants:
        name = getattr(el, 'name', None)
        if name is None:
            continue
        if name == 'tr':
            cells = el.find_all(['th', 'td'], recursive=False)
            if not cells:
                cells = el.find_all(['th', 'td'])
            if not cells:
                continue
            head = all(c.name == 'th' for c in cells)
            # 若紧邻上一行是「裸 th 头行」，且本行也是全 th → 并入（同一逻辑行被拆开）
            if (head and rows and rows[-1].get('bare')
                    and rows[-1]['head'] and not rows[-1]['cells']):
                rows.pop()
            rows.append({'cells': cells, 'head': head, 'bare': False})
        elif name == 'th' and el.parent is not None and el.parent.name != 'tr':
            if rows and rows[-1].get('bare') and rows[-1]['head']:
                rows[-1]['cells'].append(el)          # 追加到同一批裸 th
            else:
                rows.append({'cells': [el], 'head': True, 'bare': True})
    return [r for r in rows if r['cells']]


def _grid_of(rows, max_cols):
    """展开 colspan，得到二维网格 + 每行是否表头。"""
    grid = []
    for r in rows:
        vals = []
        for c in r['cells']:
            txt = _clean_text(c.get_text(' '))
            span = int(c.get('colspan', 1) or 1)
            vals.append(txt)
            vals.extend([''] * (span - 1))
        grid.append({'v': vals, 'head': r['head'],
                     'colors': [sev_color(c.get_text(' ')) for c in r['cells']
                                for _ in range(int(c.get('colspan', 1) or 1))]})
    ncol = max((len(g['v']) for g in grid), default=0)
    return grid, min(ncol, max_cols)


def _title_row_idx(grid):
    """判断首行是否是「表标题行」：只有首格有字，且次行文字格更多。"""
    if len(grid) < 2:
        return None
    def ne(v):
        return [i for i, x in enumerate(v) if x.strip()]
    a, b = ne(grid[0]['v']), ne(grid[1]['v'])
    if len(a) == 1 and a[0] == 0 and len(a[0:1]) == 1 and len(b) >= 2 \
            and not grid[0]['head']:
        return 0
    # 全 th 的标题行（colspan 合并）也按标题处理
    if len(a) == 1 and len(b) >= 2 and len(grid[0]['v']) > len(a):
        return 0
    return None


def add_table_html(doc, tbl, max_cols=22):
    """把 HTML <table> 转为 Word 表格（标题行 / 表头行分离，表头跨页重复）。"""
    rows = _table_rows(tbl)
    if not rows:
        return None
    grid, ncol = _grid_of(rows, max_cols)
    if ncol == 0:
        return None

    tidx = _title_row_idx(grid)
    has_title = tidx is not None and grid[tidx]['v'][0].strip() != ''
    body_start = 0
    if has_title:
        title_text = grid[0]['v'][0]
        body_start = 1

    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fs = 9 if ncol <= 8 else (8 if ncol <= 13 else 7)

    # —— 表标题行：整行合并、居中、加粗、浅底
    if has_title:
        r0 = t.add_row()
        for c in r0.cells:
            vmerge_none(c)
        m = r0.cells[0].merge(r0.cells[ncol - 1])
        set_cell(m, title_text, bold=True, align='center', size=fs + 0.5)
        shade(m, TITLE_FILL)

    for gi in range(body_start, len(grid)):
        g = grid[gi]
        row = t.add_row()
        is_head = g['head'] and gi == body_start
        for j in range(ncol):
            v = g['v'][j] if j < len(g['v']) else ''
            vmerge_none(row.cells[j])
            num = bool(re.fullmatch(r'[-+]?\s*[\d,\.]+%?', str(v).strip())) and str(v).strip() != ''
            al = 'center' if is_head else ('right' if num else 'left')
            col = (255, 255, 255) if is_head else g['colors'][j] if j < len(g['colors']) else None
            set_cell(row.cells[j], v, bold=is_head, color=col, align=al, size=fs)
            if is_head:
                shade(row.cells[j], HEAD_FILL)
            elif (gi - body_start) % 2 == 0:
                shade(row.cells[j], ZEBRA_FILL)
        if is_head:
            repeat_header(row)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return t


# ---------------------------------------------------------------- 卡片 → 表格
# 卡片字段白名单（避免卡片内嵌套的通用字段类被重复计入）
CARD_FIELD_SCOPE = {
    'progress-bar-wrap': ['progress-bar-label'],
}


def _field_els(card, cls):
    """在卡片内查找字段元素（允许嵌套一层），跳过被内层卡片占用的同名字段。"""
    out = []
    for el in card.select('.' + cls):
        p = el.parent
        nested = False
        while p is not None and p is not card:
            if _classes(p) & (METRIC_CARD_CLASS | set(CONTENT_CARD_CLASS)):
                nested = True
                break
            p = p.parent
        if not nested:
            out.append(el)
    return out


def _field_el(card, cls):
    els = _field_els(card, cls)
    return els[0] if els else None


def _field_text(card, cls):
    """字段文本：同名字段出现多次时全部保留（如 kpi 卡里的「全区 35.8%」小注）。"""
    txts = []
    for el in _field_els(card, cls):
        t = _clean_text(el.get_text(' '))
        if t and t not in txts:
            txts.append(t)
    return '　'.join(txts)


def _card_fields(card):
    """返回卡片内的字段类名（按 FIELD_ORDER 排序）。"""
    allowed = None
    for c in _classes(card):
        if c in CARD_FIELD_SCOPE:
            allowed = set(CARD_FIELD_SCOPE[c])
            break
    found = []
    for ch in card.find_all(True):
        for c in (ch.get('class') or []):
            if c in FIELD_LABEL and FIELD_LABEL[c] is not None:
                if allowed is not None and c not in allowed:
                    break
                if _field_els(card, c):
                    found.append(c)
                break
    rank = {k: i for i, k in enumerate(FIELD_ORDER)}
    seen, out = set(), []
    for f in found:
        if f not in seen:
            seen.add(f)
            out.append(f)
    return sorted(out, key=lambda x: rank.get(x, 999))


def _card_kind(el):
    cls = _classes(el)
    if cls & SKIP_CLASS:
        return None
    if 'strategy-card' in cls:
        return 'strategy-card'
    # 成果卡（课题成果总览）：含 num-badge / card-title
    if 'card' in cls:
        if el.select_one('.num-badge') or el.select_one('.card-title'):
            return 'result-card'
        return None
    for k in CONTENT_CARD_CLASS:
        if k in cls:
            return k
    for k in SINGLE_CARD_CLASS:
        if k in cls:
            return k
    for k in METRIC_CARD_CLASS:
        if k in cls and _card_fields(el):
            return k
    return None


def add_metric_table(doc, cards, kind):
    """一组同构指标卡 → 一张表（表头 = 字段名，每卡一行）。"""
    cols = []
    for c in cards:
        for f in _card_fields(c):
            if f not in cols:
                cols.append(f)
    rank = {k: i for i, k in enumerate(FIELD_ORDER)}
    cols.sort(key=lambda x: rank.get(x, 999))
    if not cols:
        return None
    heads = [FIELD_LABEL[c] or c for c in cols]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fs = 9 if len(cols) <= 6 else 8
    for i, h in enumerate(heads):
        set_cell(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255),
                 align='center', size=fs)
        shade(t.rows[0].cells[i], HEAD_FILL)
    repeat_header(t.rows[0])
    for ri, c in enumerate(cards):
        row = t.add_row()
        for i, f in enumerate(cols):
            raw = _field_text(c, f)
            v = _clean_text(raw)
            num = bool(re.fullmatch(r'[-+]?\s*[\d,\.]+%?\s*[\w\u4e00-\u9fff]*', v)) and v != ''
            set_cell(row.cells[i], v, align='left' if i == 0 else ('right' if num else 'left'),
                     size=fs, color=sev_color(raw))
            if ri % 2 == 1:
                shade(row.cells[i], ZEBRA_FILL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return t


def add_pair_table(doc, cards, kind, walker=None):
    """内容卡组 → 两列表（标题 | 正文），把「清单」变成表格式卡片。"""
    tsel, bsels, heads = CONTENT_CARD_CLASS[kind]
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(heads):
        set_cell(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255),
                 align='center', size=9)
        shade(t.rows[0].cells[i], HEAD_FILL)
    repeat_header(t.rows[0])
    leftovers = []
    for ri, card in enumerate(cards):
        t_el = card.select_one('.' + tsel)
        title = _clean_text(t_el.get_text(' ')) if t_el else ''
        body_parts = []
        consumed = {id(t_el)} if t_el is not None else set()
        for b in bsels:
            el = card.select_one('.' + b)
            if el is None:
                continue
            consumed.add(id(el))
            txt = _clean_text(el.get_text(' '))
            if txt:
                body_parts.append(txt)
        # field-cat：正文是 <li> 清单
        if not body_parts:
            items = [_clean_text(li.get_text(' ')) for li in card.select('li')]
            body_parts = ['、'.join(x for x in items if x)] if items else []
        if not title and not body_parts:
            continue
        if not title:
            title = body_parts.pop(0)
        row = t.add_row()
        set_cell(row.cells[0], title, bold=True, align='left', size=9)
        shade(row.cells[0], CARD_TITLE_FILL)
        set_cell(row.cells[1], '　'.join(body_parts), align='left', size=9)
        if ri % 2 == 1:
            shade(row.cells[1], ZEBRA_FILL)
        for ch in card.find_all(recursive=False):
            if id(ch) not in consumed and _clean_text(ch.get_text(' ')):
                leftovers.append(ch)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if leftovers and walker is not None:
        for el in leftovers:
            if el.name == 'table':
                add_table_html(doc, el)
            else:
                walker(el)
    return t


def add_single_card(doc, card, kind, walker=None):
    """单张（标题 + 正文）卡 → 加粗标题 + 正文段，剩余块继续递归。"""
    tsel, bsels = SINGLE_CARD_CLASS[kind]
    consumed = set()
    t_el = card.select_one('.' + tsel)
    if t_el:
        consumed.add(id(t_el))
        txt = _clean_text(t_el.get_text(' '))
        if txt:
            add_card_title(doc, txt)
    for b in bsels:
        el = card.select_one('.' + b)
        if el is None:
            continue
        consumed.add(id(el))
        txt = _clean_text(el.get_text(' '))
        if txt:
            add_body(doc, txt, indent=False, size=10.5)
    for ch in card.find_all(recursive=False):
        if id(ch) in consumed:
            continue
        if el_empty(ch):
            continue
        if walker is not None:
            walker(ch)


def el_empty(el):
    return not _clean_text(el.get_text(' ')) and el.name != 'table'


def add_callout(doc, text, fill=CALLOUT_FILL):
    """提示框 → 单格浅底表（保留视觉上的“强调块”）。"""
    if not text:
        return
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    set_cell(cell, text, align='left', size=10)
    shade(cell, fill)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)


def add_mono_block(doc, text):
    """预格式化文本（ASCII 表格 / 代码）→ 单格灰底等宽块，保留换行与对齐。"""
    if not text:
        return
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    for line in text.split('\n'):
        p = t.rows[0].cells[0].paragraphs[0] if line == text.split('\n')[0] \
            else t.rows[0].cells[0].add_paragraph()
        r = p.add_run(line.rstrip())
        r.font.size = Pt(7)
        r.font.name = 'Consolas'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    shade(t.rows[0].cells[0], MONO_FILL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)


def add_bar_chart_table(doc, chart):
    """CSS 柱状图（数据存在 title 属性里）→ 数据表。"""
    bars = chart.select('.bar[title]')
    if not bars:
        return None
    # 取全部 title，形如「2024年: 305件」
    series, rows = [], []
    for g in chart.select('.bar-group'):
        lab = g.select_one('.bar-label')
        row = [(_clean_text(lab.get_text(' ')) if lab else '')]
        for b in g.select('.bar[title]'):
            ti = b.get('title') or ''
            name, _, val = ti.partition(':')
            name = _clean_text(name)
            val = _clean_text(val)
            if name and name not in series:
                series.append(name)
            row.append(val)
        rows.append(row)
    if not rows:
        return None
    if not series:
        series = ['数值']
    t = doc.add_table(rows=1, cols=1 + len(series))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(t.rows[0].cells[0], '月份', bold=True, color=(255, 255, 255),
             align='center', size=9)
    shade(t.rows[0].cells[0], HEAD_FILL)
    for i, h in enumerate(series):
        set_cell(t.rows[0].cells[i + 1], h, bold=True, color=(255, 255, 255),
                 align='center', size=9)
        shade(t.rows[0].cells[i + 1], HEAD_FILL)
    repeat_header(t.rows[0])
    for ri, row in enumerate(rows):
        r = t.add_row()
        for i, v in enumerate(row[:1 + len(series)]):
            set_cell(r.cells[i], v, align='center', size=9)
            if ri % 2 == 1:
                shade(r.cells[i], ZEBRA_FILL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return t



def add_result_card_table(doc, cards):
    """课题成果总览的成果卡 → 一张 3 列表。"""
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(['成果', '内容', '标签 / 附件']):
        set_cell(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255),
                 align='center', size=9)
        shade(t.rows[0].cells[i], HEAD_FILL)
    repeat_header(t.rows[0])
    for ri, c in enumerate(cards):
        row = t.add_row()
        badge = c.select_one('.num-badge')
        title = c.select_one('.card-title')
        first = ' '.join(x for x in [_clean_text(badge.get_text(' ')) if badge else '',
                                     _clean_text(title.get_text(' ')) if title else ''] if x)
        desc = c.select_one('.card-desc')
        tags = c.select('.card-tag')
        body = _clean_text(desc.get_text(' ')) if desc else ''
        if tags:
            body = (body + '　［' + '、'.join(_clean_text(x.get_text(' ')) for x in tags) + '］').strip()
        att = c.select_one('.card-attach')
        attach = _clean_text(att.get_text(' ')) if att else ''
        for i, v in enumerate([first, body, attach]):
            set_cell(row.cells[i], v, bold=(i == 0), align='left', size=9)
            if ri % 2 == 1:
                shade(row.cells[i], ZEBRA_FILL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return t


def add_flow_table(doc, chart):
    """流程图 → 「步骤 | 说明」两列表（Word 里比箭头串更易读）。"""
    nodes = chart.select('.flow-node')
    if not nodes:
        return None
    rows = []
    for nd in nodes:
        cls = _classes(nd)
        txt = _clean_text(nd.get_text(' '))
        if not txt:
            continue
        kind = ('起止' if (cls & {'start', 'end'}) else
                '判断' if 'decision' in cls else '步骤')
        rows.append((kind, txt))
    if not rows:
        return None
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(['环节', '内容']):
        set_cell(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255),
                 align='center', size=9)
        shade(t.rows[0].cells[i], HEAD_FILL)
    repeat_header(t.rows[0])
    seen = set()
    for ri, (kind, txt) in enumerate(rows):
        if txt in seen:
            continue
        seen.add(txt)
        row = t.add_row()
        set_cell(row.cells[0], kind, bold=True, align='center', size=9)
        shade(row.cells[0], CARD_TITLE_FILL)
        set_cell(row.cells[1], txt, align='left', size=9)
        if ri % 2 == 1:
            shade(row.cells[1], ZEBRA_FILL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return t


# ---- 内联样式卡片（无 class，仅靠 style 定义外观）--------------------
FS_RE = re.compile(r'font-size\s*:\s*([\d.]+)px')
BOLD_RE = re.compile(r'font-weight\s*:\s*(bold|[6-9]\d\d)')
LEGEND_RE = re.compile(r'^\s*(红|橙|黄|绿|蓝|紫|灰|青|深色|浅色)色?\s*[:：]')
HIDE_RE = re.compile(r'display\s*:\s*none')
# display:none 的「例外」：这些区块只是被 tab／筛选器默认收起，内容本身有效，Word 里必须带上
# （反例：成果总览的 #tools / #consulting / #docs 是「待推进」占位，不应进交付物 → 不列入例外）
HIDE_ALLOW_CLASS = {'year-content'}       # 数据分析月报合集的 2025/2026 月份卡网格
HIDE_ALLOW_ID = {'saViewCommunity'}       # 治理效果追踪分析的「小区视图」面板


def _style(el):
    return el.get('style') or ''


def _is_value(t):
    return bool(re.fullmatch(r'[-+]?[\d,\.]+%?(件|个|户|元|人|天|条|分)?', t or '')) \
        or (len(t or '') <= 8 and any(c.isdigit() for c in t or ''))


def _has_block_kid(el):
    return any(_is_block(k) and not _skip(k)
               for k in el.children if getattr(k, 'name', None))


def _is_styled_card(el):
    """内联样式写的「指标卡」：无 class，带底色/左边框，子元素都是纯文本。"""
    if el.name != 'div' or _classes(el) or el.get('id') or _skip(el):
        return False
    st = _style(el)
    if 'background' not in st and 'border-left' not in st:
        return False
    kids = [k for k in el.children if getattr(k, 'name', None) and not _skip(k)]
    if not kids:
        return False
    return not any(_has_block_kid(k) for k in kids)


def _is_styled_panel(el):
    """内联样式写的「左右分栏板块」：带底色 + 有块级子元素（标题 + 清单）。"""
    if el.name != 'div' or _classes(el) or el.get('id') or _skip(el):
        return False
    st = _style(el)
    if 'background' not in st and 'border-left' not in st:
        return False
    return bool([k for k in el.children if getattr(k, 'name', None) and not _skip(k)])


def _is_styled_title(el):
    """内联样式的小标题（font-weight 加粗 + 短文本、非纯数字）。"""
    if el.name not in ('div', 'p', 'span') or _classes(el) or _skip(el):
        return False
    if not BOLD_RE.search(_style(el)):
        return False
    txt = _clean_text(el.get_text(' '))
    return bool(txt) and len(txt) <= 80 and not _is_value(txt)


def _card_parts(card):
    """把内联样式卡片拆成 [(字号, 文本, 是否数值)]。"""
    parts = []
    for k in card.children:
        if getattr(k, 'name', None) is None or k.name == 'br' or _skip(k):
            continue
        txt = _clean_text(k.get_text(' '))
        if not txt:
            continue
        m = FS_RE.search(_style(k))
        fs = float(m.group(1)) if m else 12.0
        parts.append((fs, txt, _is_value(txt)))
    return parts


def _styled_table(doc, heads, rows, aligns=None):
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(heads):
        set_cell(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255),
                 align='center', size=9)
        shade(t.rows[0].cells[i], HEAD_FILL)
    repeat_header(t.rows[0])
    aligns = aligns or (['left'] * len(heads))
    for ri, rowv in enumerate(rows):
        row = t.add_row()
        for i, v in enumerate(rowv):
            set_cell(row.cells[i], v, bold=(i == 0), align=aligns[i], size=9,
                     color=sev_color(v))
            if ri % 2 == 1:
                shade(row.cells[i], ZEBRA_FILL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return t


def add_styled_group(doc, cards):
    """网格里的内联样式卡片组 → 指标表 / 要点表（把「清单」变成卡片表）。"""
    parsed = [_card_parts(c) for c in cards]
    parsed = [p for p in parsed if p]
    if not parsed:
        return None
    numeric = all(any(x[2] for x in parts) for parts in parsed)
    rows = []
    if numeric:
        for parts in parsed:
            order = sorted(range(len(parts)), key=lambda i: -parts[i][0])
            vi = next((i for i in order if parts[i][2]), order[0])
            rest = [i for i in order if i != vi]
            label = parts[rest[0]][1] if rest else ''
            note = '；'.join(parts[i][1] for i in rest[1:])
            rows.append((label, parts[vi][1], note))
        return _styled_table(doc, ['项目', '数值', '说明'], rows,
                             aligns=['left', 'right', 'left'])
    for parts in parsed:
        if len(parts) == 1:
            rows.append((parts[0][1], ''))
        else:
            rows.append((parts[0][1], '　'.join(x[1] for x in parts[1:])))
    return _styled_table(doc, ['要点', '说明'], rows, aligns=['left', 'left'])


def add_panel_table(doc, panels):
    """左右分栏的「预警 / 成效」板块 → 两列表。"""
    rows = []
    for pn in panels:
        title = ''
        items = []
        for k in pn.children:
            if getattr(k, 'name', None) is None or _skip(k):
                continue
            if not title and (BOLD_RE.search(_style(k)) or k.name in ('h3', 'h4', 'h5')):
                title = _clean_text(k.get_text(' '))
                continue
            lis = k.select('li') if k.name in ('ul', 'ol') else []
            if lis:
                items.extend(_clean_text(li.get_text(' ')) for li in lis)
            else:
                txt = _clean_text(k.get_text(' '))
                if txt:
                    items.append(txt)
        if not title:
            title = next((_clean_text(li.get_text(' ')) for li in pn.select('li')), '')
            items = items[1:] if items else items
        items = [x for x in items if x]
        if title or items:
            rows.append((title or '—', '；'.join(items)))
    if not rows:
        return None
    return _styled_table(doc, ['专项板块', '明细'], rows, aligns=['left', 'left'])


def _ancestor_chain(el, n=5):
    out, cur = [], el
    for _ in range(n):
        if cur is None or cur.name is None:
            break
        cls = '.'.join(cur.get('class') or [])
        eid = cur.get('id')
        st = (cur.get('style') or '')[:40]
        tag = cur.name + ('.' + cls if cls else '')
        if eid:
            tag += '#' + eid
        if st:
            tag += '[' + st + ']'
        out.append(tag)
        cur = cur.parent
    return ' ← '.join(out)


def add_strategy_card(doc, card):
    """去重方法论的策略卡 → 小标题 + （判定条件 / 识别逻辑 / 识别结果）两列表。"""
    badge = card.select_one('.strategy-badge')
    title = card.select_one('.strategy-title')
    head = ' '.join(x for x in [_clean_text(badge.get_text(' ')) if badge else '',
                                _clean_text(title.get_text(' ')) if title else ''] if x)
    if head:
        add_card_title(doc, head, size=11.5)
    desc = card.select_one('.strategy-desc')
    if desc:
        txt = _clean_text(desc.get_text(' '))
        if txt:
            add_body(doc, txt, indent=False, size=10.5)
    pairs = []
    for box in card.select('.criteria-box'):
        lb = box.select_one('.criteria-label')
        vl = box.select_one('.criteria-value')
        if vl is None:
            continue
        pairs.append((_clean_text(lb.get_text(' ')) if lb else '', _clean_text(vl.get_text(' '))))
    if pairs:
        t = doc.add_table(rows=0, cols=2)
        t.style = 'Table Grid'
        for i, (lb, vl) in enumerate(pairs):
            row = t.add_row()
            set_cell(row.cells[0], lb, bold=True, align='center', size=9)
            shade(row.cells[0], CARD_TITLE_FILL)
            set_cell(row.cells[1], vl, align='left', size=9)
            if i % 2 == 1:
                shade(row.cells[1], ZEBRA_FILL)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
    # 其余块
    for ch in card.find_all(recursive=False):
        cl = set(ch.get('class') or [])
        if cl & {'strategy-header', 'strategy-desc', 'criteria-box'}:
            continue
        txt = _clean_text(ch.get_text(' '))
        if txt:
            add_body(doc, txt, indent=False, size=10.5)


# ---------------------------------------------------------------- 解析
def _classes(el):
    return set(el.get('class') or [])


def _skip(el):
    if el.name in SKIP_TAG:
        return True
    if _classes(el) & SKIP_CLASS:
        return True
    if el.get('aria-hidden') == 'true':
        return True
    if (HIDE_RE.search(el.get('style') or '')
            and not (_classes(el) & HIDE_ALLOW_CLASS)
            and el.get('id') not in HIDE_ALLOW_ID):
        return True
    return False


def _is_block(el):
    return el.name in BLOCK_TAG


def sev_color(s):
    """源报告用彩色圆点表达严重度；返回对应文字颜色。"""
    for ch, rgb in SEV_COLOR.items():
        if ch in (s or ''):
            return rgb
    return None


def _clean_text(s):
    s = re.sub(r'\s+', ' ', s or '').strip()
    s = s.replace('\u200b', '')
    # 去掉 emoji（报告中用作装饰/等级标识，Word 里改用文字颜色 / 文字分级表达）
    # ★☆（U+2605/2606）是治理难度评级符号，保留
    s = re.sub(r'[\U0001F000-\U0001FAFF\uFE0F\u2700-\u27BF\u2B00-\u2BFF'
               r'\u2300-\u23FF\uE000-\uF8FF]', '', s)
    s = re.sub(r'[\u2600-\u26FF]', lambda m: m.group(0) if m.group(0) in '★☆' else '', s)
    # 修掉标签分隔造成的多余空格：数字与单位/%之间
    s = re.sub(r'(?<=[\d])\s+(?=[%‰件个元人天户条倍])', '', s)
    s = re.sub(r'\s+([，。、；：）】》%‰])', r'\1', s)
    s = re.sub(r'([（【《])\s+', r'\1', s)
    return s.strip()


def _heading_level(el, base):
    m = {'h1': 1, 'h2': 2, 'h3': 3, 'h4': 4, 'h5': 4, 'h6': 4}
    lv = m.get(el.name, 2)
    return min(9, max(1, lv + base - 1))


def load_figmap(manifest, stem):
    """返回 {容器id: 绝对路径}"""
    figs = manifest.get(stem, [])
    return {f['id']: os.path.join(ROOT, f['file']) for f in figs}


def load_figlist(manifest, stem):
    return list(manifest.get(stem, []))


def convert_report(doc, html_path, figmap, base=1, chapter_title=None,
                   skip_heading1=True, on_image=None, fig_prefix=None,
                   stats=None):
    """把一份报告写入 doc。base=标题层级基准（1 表示报告 h1→Heading1）。"""
    html = open(html_path, encoding='utf-8').read()
    soup = BeautifulSoup(html, 'lxml')
    body = soup.body or soup
    for t in body.find_all(SKIP_TAG):
        t.decompose()

    fig_used = []
    fig_no = [0]
    if stats is None:
        stats = {}
    stats.setdefault('tables', 0)
    stats.setdefault('cards', 0)
    stats.setdefault('title_rows', 0)

    def render_one(child):
        cls = _classes(child)
        # 图表容器 → 插图（各报告容器 class 不统一：chart-box / chart）
        if child.get('id') and ('chart-box' in cls or 'chart' in cls
                                or any(c.startswith('chart-') for c in cls)):
            cid = child.get('id')
            fp = figmap.get(cid)
            title = child.get('data-title') or ''
            if fp:
                fig_no[0] += 1
                cap = '图 %s-%d%s' % (fig_prefix, fig_no[0], ('　' + title) if title else '') \
                    if fig_prefix else (title or None)
                add_image(doc, fp, caption=cap)
                fig_used.append(cid)
                if on_image:
                    on_image(cid)
            return
        if child.name == 'table':
            add_table_html(doc, child)
            stats['tables'] += 1
            return
        if child.name in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            txt = _clean_text(child.get_text(' '))
            if not txt:
                return
            lv = _heading_level(child, base)
            if child.name == 'h1' and skip_heading1:
                return
            doc.add_heading(txt, level=lv)
            return
        if 'section-title' in cls:
            txt = _clean_text(child.get_text(' '))
            if txt:
                doc.add_heading(txt, level=min(9, base + 1))
            return
        # 分区头（成果总览的「一 / 核心硬成果 / 3项已完成」）
        if 'cat-header' in cls:
            t1 = child.select_one('.cat-title')
            t2 = child.select_one('.cat-count')
            s1 = _clean_text(t1.get_text(' ')) if t1 else ''
            s2 = _clean_text(t2.get_text(' ')) if t2 else ''
            if s1:
                doc.add_heading(s1, level=min(9, base + 1))
            if s2:
                add_body(doc, s2, indent=False, size=10.5)
            return
        # 关联模型的主表节点
        if 'model-center' in cls:
            txt = _clean_text(child.get_text(' '))
            if txt:
                add_card_title(doc, txt)
            return
        if 'section-desc' in cls or 'sub-title' in cls:
            txt = _clean_text(child.get_text(' '))
            if txt:
                add_body(doc, txt, indent=False, size=10.5)
            return
        # 报告页眉（subtitle / meta）→ 居中小字说明 + 其余统计卡照常渲染
        if child.name == 'header' or 'header' in cls:
            sub = child.select_one('.subtitle')
            meta = child.select_one('.meta')
            for el in (sub, meta):
                if el is None:
                    continue
                txt = _clean_text(el.get_text(' '))
                if not txt:
                    continue
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt)
                r.font.size = Pt(10 if el is sub else 9)
                r.font.name = CN_FONT
                r._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
                if el is meta:
                    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                p.paragraph_format.space_after = Pt(3)
            if sub is not None:
                sub.decompose()
            if meta is not None:
                meta.decompose()
            walk(child)
            return
        # 预格式化文本（ASCII 表格 / 代码块）
        if cls & MONO_CLASS:
            add_mono_block(doc, child.get_text('\n'))
            return
        # 提示框（浅底强调块）
        if (cls & CALLOUT_CLASS) and child.name == 'div':
            add_callout(doc, _clean_text(child.get_text(' ')))
            return
        # 关键词云 → 归并成一段
        if cls & KEYWORD_CLASS:
            words = [_clean_text(x.get_text(' ')) for x in child.find_all('span')]
            words = [w for w in words if w]
            if words:
                add_body(doc, '关键词：' + '、'.join(words), indent=False, size=10.5)
            return
        # 小标题 / 脚注 / 图注
        if cls & CARD_TITLE_CLASS:
            txt = _clean_text(child.get_text(' '))
            if txt:
                add_card_title(doc, txt)
            return
        if cls & CAPTION_CLASS:
            txt = _clean_text(child.get_text(' '))
            if txt:
                add_caption(doc, txt)
            return
        # CSS 柱状图 → 数据表
        if cls & BARCHART_CLASS:
            add_bar_chart_table(doc, child)
            stats['barchart'] = stats.get('barchart', 0) + 1
            return
        # 图表容器（无导出图时退化为数据表）
        if ('chart-box' in cls or 'chart-body' in cls or 'chart-wrap' in cls
                or 'chart-area' in cls or any(c.endswith('-chart') for c in cls)):
            add_bar_chart_table(doc, child)
        # 地图容器：无图时保留标题 + 图注
        if 'map-container' in cls:
            for sel in ('map-title', 'map-caption'):
                el = child.select_one('.' + sel)
                if el is None:
                    continue
                txt = _clean_text(el.get_text(' '))
                if not txt:
                    continue
                if sel == 'map-title':
                    add_card_title(doc, txt)
                else:
                    add_caption(doc, txt)
            return
        # 流程图 → 节点 + 说明两列表
        if 'flowchart' in cls:
            add_flow_table(doc, child)
            return
        # 专题头部（编号 + 标题 + 副标题）→ 标题 + 说明
        if 'topic-header' in cls:
            num = child.select_one('.topic-num')
            h = child.find(['h1', 'h2', 'h3', 'h4'])
            ps = child.find('p')
            head = ' '.join(x for x in [_clean_text(num.get_text(' ')) if num else '',
                                        _clean_text(h.get_text(' ')) if h else ''] if x)
            if head:
                doc.add_heading(head, level=min(9, base + 1))
            if ps is not None:
                txt = _clean_text(ps.get_text(' '))
                if txt:
                    add_body(doc, txt, indent=False, size=10.5)
            return
        # 报告页眉以外的图例说明（「绿色：2025年同比下降」）→ 丢弃
        first_txt = _clean_text(child.get_text(' '))
        if first_txt and len(first_txt) <= 40 and LEGEND_RE.match(first_txt):
            return
        # 内联样式的网格容器（指标卡组 / 左右板块组）
        if 'display:grid' in _style(child):
            kids = [k for k in child.children
                    if getattr(k, 'name', None) and not _skip(k)]
            if kids and all(_is_styled_card(k) for k in kids):
                add_styled_group(doc, kids)
                return
            if kids and all(_is_styled_panel(k) for k in kids):
                add_panel_table(doc, kids)
                return
        # 内联样式的小标题
        if _is_styled_title(child):
            add_card_title(doc, _clean_text(child.get_text(' ')))
            return
        # 孤立的数据问题 / 解决方案块 → 提示框
        if cls & {'issue-solution', 'issue-desc', 'solution'}:
            add_callout(doc, first_txt)
            return
        # 无 class、仅靠内联样式着色的提示块 → 提示框
        if (child.name == 'div' and not cls and not child.get('id')
                and not _has_block_kid(child)
                and ('border-left' in _style(child) or 'background' in _style(child))
                and first_txt):
            add_callout(doc, first_txt)
            return
        if child.name == 'li':
            txt = _clean_text(child.get_text(' '))
            if txt:
                add_bullet(doc, txt)
            return
        if child.name == 'p':
            txt = _clean_text(child.get_text(' '))
            if txt:
                add_body(doc, txt)
            return
        if cls & ATOMIC_CLASS:
            txt = _clean_text(child.get_text(' '))
            if txt:
                add_body(doc, txt, indent=False, size=10.5)
            return
        # 块级容器：有块级子元素则递归，否则作为文本叶子
        kids = [k for k in child.children if getattr(k, 'name', None)]
        has_block = any(_is_block(k) and not _skip(k) for k in kids)
        if has_block:
            walk(child, 0)
        else:
            txt = _clean_text(child.get_text(' '))
            if txt:
                key = ','.join(sorted(cls)) or child.name
                lv = stats.setdefault('leaves', {})
                lv[key] = lv.get(key, 0) + 1
                smp = stats.setdefault('leaf_samples', {})
                ch = _ancestor_chain(child)
                d = smp.setdefault(key, {})
                if ch in d:
                    d[ch][0] += 1
                elif len(d) < 60:
                    d[ch] = [1, txt[:70]]
                add_body(doc, txt)

    def walk_list(kids, owner):
        """按卡片规则遍历一组兄弟元素（owner 仅用于定位行内 run 的归属）。"""
        i = 0
        inline_buf = []

        def flush_inline():
            if not inline_buf:
                return
            txt = _clean_text(' '.join(x.get_text(' ') for x in inline_buf))
            bold_only = all(x.name in ('strong', 'b') for x in inline_buf)
            key = 'INLINE:%s' % owner.name
            chain = _ancestor_chain(owner)
            inline_buf.clear()
            if not txt:
                return
            # 纯加粗短句（如「针对性建议：」）→ 作小标题，避免出现孤立清单项
            if bold_only and len(txt) <= 40:
                add_card_title(doc, txt)
                return
            lv = stats.setdefault('leaves', {})
            lv[key] = lv.get(key, 0) + 1
            d = stats.setdefault('leaf_samples', {}).setdefault(key, {})
            if chain in d:
                d[chain][0] += 1
            elif len(d) < 60:
                d[chain] = [1, txt[:70]]
            add_body(doc, txt)

        while i < len(kids):
            child = kids[i]
            if _skip(child):
                i += 1
                continue
            # 行内元素：与相邻行内兄弟合并为一段，避免句子被拆碎
            if (child.name != 'table' and not _has_block_kid(child)
                    and (child.name in INLINE_TAG or not _is_block(child))):
                if child.name == 'br':
                    i += 1
                    continue
                inline_buf.append(child)
                i += 1
                continue
            flush_inline()
            kind = _card_kind(child)
            if kind:
                group = [child]
                j = i + 1
                while j < len(kids):
                    nx = kids[j]
                    if _skip(nx):
                        j += 1
                        continue
                    if _card_kind(nx) == kind:
                        group.append(nx)
                        j += 1
                    else:
                        break
                if kind == 'strategy-card':
                    for c in group:
                        add_strategy_card(doc, c)
                        stats['cards'] += 1
                elif kind == 'result-card':
                    add_result_card_table(doc, group)
                    stats['cards'] += len(group)
                elif kind in CONTENT_CARD_CLASS:
                    add_pair_table(doc, group, kind, walker=walk)
                    stats['cards'] += len(group)
                elif kind in SINGLE_CARD_CLASS:
                    for c in group:
                        add_single_card(doc, c, kind, walker=walk)
                        stats['cards'] += 1
                else:
                    add_metric_table(doc, group, kind)
                    stats['cards'] += len(group)
                i = j
                continue
            render_one(child)
            i += 1
        flush_inline()

    def walk(el, depth=0):
        walk_list([k for k in el.children if getattr(k, 'name', None) is not None], el)

    walk(body)
    return fig_used

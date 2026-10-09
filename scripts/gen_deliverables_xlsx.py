#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 4 份 Excel 全量清单交付附件 → 交付物/

  附件1_高风险小区评分全量清单.xlsx        772 小区 × 22 指标
  附件2_2026年1-8月投诉去重建议清单.xlsx   8 个月 × 6 张表
  附件3_词云清单.xlsx                     13 街镇 × 60 词（Top24 展示 + 补位候选）
  附件4_交付物清单.xlsx                    交付物目录全量文件目录（含过程成果报告）

（v2.17.0 起取消原「附件3 热线与纳统小区全量匹配表」「附件4 纳统匹配候选与未匹配清单」；
  v2.17.4 起装帧素材（封面 / 书脊 / 封底 / 全封展开图与封面候选底图）不纳入交付物清单）

用法：
  PYTHONPATH=~/.workbuddy/binaries/python/vendor /usr/bin/python3 scripts/gen_deliverables_xlsx.py
"""
import json
import os
import sys

import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, '交付物')

HDR_FILL = PatternFill('solid', fgColor='1A3A5F')
HDR_FONT = Font(color='FFFFFF', bold=True, size=10, name='微软雅黑')
TITLE_FONT = Font(bold=True, size=12, color='1A3A5F', name='微软雅黑')
BODY_FONT = Font(size=10, name='微软雅黑')
THIN = Side(style='thin', color='CCCCCC')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)

TBL_NAMES = ['识别策略汇总', '高频 20 小区', '文本相似典型案例', '类别分布', '街道分布', '处置建议清单']


def read_xls(fp, **kw):
    """统一按字符串读取（keep_default_na=False），避免长数字 ID 被读成科学计数法。"""
    return pd.read_excel(fp, dtype=str, keep_default_na=False, **kw)


def write_block(ws, r, title, df):
    """在 ws 的第 r 行起写入「标题 + 表头 + 数据」，返回下一块起始行。"""
    if title:
        c = ws.cell(row=r, column=1, value=title)
        c.font = TITLE_FONT
        r += 1
    ncol = max(len(df.columns), 1)
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=r, column=j, value=str(col))
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, CENTER, BORDER
    ws.row_dimensions[r].height = 22
    r += 1
    for _, row in df.iterrows():
        for j, col in enumerate(df.columns, start=1):
            v = row[col]
            if pd.isna(v):
                v = ''
            c = ws.cell(row=r, column=j, value=(v.item() if hasattr(v, 'item') else v))
            c.font, c.alignment, c.border = BODY_FONT, CENTER, BORDER
        r += 1
    return r + 1, ncol


def autofit(ws, ncol, widths=None):
    for j in range(1, ncol + 1):
        L = get_column_letter(j)
        if widths and j <= len(widths):
            ws.column_dimensions[L].width = widths[j - 1]
        else:
            ws.column_dimensions[L].width = 14


# ---------- 附件1：高风险小区评分全量清单 ----------
def deliver_risk():
    src = os.path.join(ROOT, '高风险小区预警清单.xlsx')
    df = read_xls(src)
    wb = Workbook()
    ws = wb.active
    ws.title = '风险评分全量清单'
    ws['A1'] = '徐汇区高风险小区评分全量清单（参评 %d 个纳统小区 · 数据截至 2026年8月）' % len(df)
    ws['A1'].font = TITLE_FONT
    nxt, ncol = write_block(ws, 3, None, df)
    widths = [18, 10, 26] + [10] * (ncol - 3)
    autofit(ws, ncol, widths)
    ws.freeze_panes = 'A4'
    fp = os.path.join(OUT, '附件1_高风险小区评分全量清单.xlsx')
    wb.save(fp)
    return fp, len(df)


# ---------- 附件2：2026年1-8月投诉去重建议清单 ----------
def deliver_dedup():
    html = open(os.path.join(ROOT, '2026年1-8月投诉去重建议清单.html'), encoding='utf-8').read()
    soup = BeautifulSoup(html, 'lxml')
    wb = Workbook()
    wb.remove(wb.active)
    total_rows = 0
    for p in soup.select('[id^=panel-]'):
        pid = p.get('id', '')
        label = pid.replace('panel-m', '').strip() + '月'
        tbs = p.find_all('table')
        if not tbs:
            continue
        ws = wb.create_sheet('%s' % label)
        ws['A1'] = '2026年%s 投诉去重建议清单' % label
        ws['A1'].font = TITLE_FONT
        r = 3
        maxc = 1
        for i, t in enumerate(tbs):
            rows = []
            for tr in t.find_all('tr'):
                cells = [c.get_text(' ', strip=True) for c in tr.find_all(['th', 'td'])]
                if any(x for x in cells):
                    rows.append(cells)
            if not rows:
                continue
            head = rows[0]
            body = rows[1:]
            ndf = pd.DataFrame(body, columns=head) if body else pd.DataFrame(columns=head)
            nm = TBL_NAMES[i] if i < len(TBL_NAMES) else '表%d' % (i + 1)
            r, nc = write_block(ws, r, nm, ndf)
            maxc = max(maxc, nc)
            total_rows += len(ndf)
        autofit(ws, maxc)
        ws.freeze_panes = 'A3'
    fp = os.path.join(OUT, '附件2_2026年1-8月投诉去重建议清单.xlsx')
    wb.save(fp)
    return fp, total_rows


# ---------- 附件3：词云清单 ----------
def deliver_wordcloud():
    src = os.path.join(ROOT, 'scripts', 'wordcloud_terms.json')
    d = json.load(open(src, encoding='utf-8'))
    meta, streets = d['meta'], d['streets']
    topn = meta['topn_shown']
    rows = []
    for st in streets:
        for it in streets[st]:
            rows.append([st, it['rank'], it['word'], it['cnt'],
                         'Top%d 展示' % topn if it['rank'] <= topn else '补位候选'])
    df = pd.DataFrame(rows, columns=['街镇', '排名', '词条', '出现次数', '状态'])
    wb = Workbook()
    ws = wb.active
    ws.title = '词云词表'
    ws['A1'] = ('徐汇区街镇诉求主题词云词表（%d 个街镇 · %d 个候选词 · 每镇 Top%d 展示）'
                % (meta['streets'], meta['unique_words'], topn))
    ws['A1'].font = TITLE_FONT
    nxt, ncol = write_block(ws, 3, None, df)
    autofit(ws, ncol, [12, 8, 16, 10, 12])
    ws.freeze_panes = 'A4'

    ws2 = wb.create_sheet('口径说明')
    ws2['A1'] = '词云词表口径'
    ws2['A1'].font = TITLE_FONT
    info = pd.DataFrame([
        ['统计周期', meta['scope']],
        ['统计单元', meta['unit']],
        ['展示口径', '每街镇按词频降序取前 %d 个词上词云；第 %d 名以后列为「补位候选」'
                     '（删词后自动补位）' % (topn, topn + 1)],
        ['候选池', '每街镇 Top%d，全区合计 %d 个候选取词' % (meta['cand_n'], meta['unique_words'])],
        ['数据来源', meta['source']],
        ['说明', '停用词表仅作用于关键词提取，不影响任何投诉量统计数字。'],
    ], columns=['项目', '说明'])
    write_block(ws2, 3, None, info)
    autofit(ws2, 2, [14, 76])

    fp = os.path.join(OUT, '附件3_词云清单.xlsx')
    wb.save(fp)
    return fp, len(df)


# ---------- 附件4：交付物清单 ----------
MANIFEST_SHEET = '交付物清单'
CATE_ORDER = ['主交付物', '电子附件', '过程成果报告', '历史版本']


def _manifest_rows():
    """扫描 交付物/ 全量文件 → [类别, 名称, 相对路径, 大小, 说明]"""
    rows = []

    def size_of(p):
        try:
            return '%.0f KB' % (os.path.getsize(p) / 1024)
        except OSError:
            return '—'

    def walk(d, prefix, cate, note):
        if not os.path.isdir(d):
            return
        for name in sorted(os.listdir(d)):
            if name.startswith('.'):
                continue
            p = os.path.join(d, name)
            rel = os.path.relpath(p, OUT)
            if os.path.isdir(p):
                walk(p, prefix + name + '/', cate, note)
            else:
                rows.append([cate, name, rel, size_of(p), note])

    for name in sorted(os.listdir(OUT)):
        if name.startswith('.'):
            continue
        p = os.path.join(OUT, name)
        rel = os.path.relpath(p, OUT)
        if os.path.isdir(p):
            if name == '_archive':       # 历史版本留档目录（不计入过程成果报告）
                walk(p, '', '历史版本', '过程版本留档，非最终交付版本')
            elif name == '装订打印封面':
                # v2.17.4：成书装帧素材不纳入交付物清单
                continue
            else:
                cate = '过程成果报告'
                note = {'月度分析报告': '月度滚动分析成果（19 份）',
                        '年度分析报告': '年度分析成果 · 街镇 13 份 + 集团 2 份'}.get(name, '过程成果')
                walk(p, '', cate, note)
            continue
        if name.endswith('.docx'):
            if '人工修改版' in name or '（一期）' in name:
                rows.append(['历史版本', name, rel, size_of(p), '过程版本留档，非最终交付版本'])
            else:
                rows.append(['主交付物', name, rel, size_of(p),
                             '课题成果汇编（正文，含图表与附录）'])
        elif name.endswith('.xlsx'):
            note = {'附件1': '高风险小区评分全量清单（772 个参评小区）',
                    '附件2': '2026 年 1-8 月投诉去重建议清单（逐月明细）',
                    '附件3': '13 街镇词云词表（Top24 展示 + 补位候选）',
                    '附件4': '交付物清单（本表）'}.get(name[:3], '')
            rows.append(['电子附件', name, rel, size_of(p), note])
        elif name.endswith('.pdf'):
            # v2.17.5：PDF 版成果汇编（由 scripts/export_pdf.py 从 docx 导出）
            rows.append(['主交付物', name, rel, size_of(p),
                         '课题成果汇编 PDF 版（打印与分发用，与 Word 版同源）'])
        elif name.endswith(('.png', '.jpg', '.jpeg')):
            # v2.17.4：封面候选底图等图片素材不纳入交付物清单
            continue
        else:
            rows.append(['电子附件', name, rel, size_of(p), ''])
    order = {c: i for i, c in enumerate(CATE_ORDER)}
    # 本表自身：首次运行时尚未落盘，按预测补一行，保证清单不随运行次序变化
    self_name = '附件4_交付物清单.xlsx'
    if not any(r[1] == self_name for r in rows):
        rows.append(['电子附件', self_name, self_name, '—', '交付物清单（本表）'])
    rows.sort(key=lambda r: (order.get(r[0], 99), r[2]))
    return rows


def deliver_manifest():
    rows = _manifest_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = '交付物清单'
    ws['A1'] = '徐汇区住宅小区治理与物业服务质量提升探索项目课题（一期）交付物清单'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = '共 %d 个文件；%s' % (len(rows), ' · '.join(
        '%s %d 份' % (c, sum(1 for r in rows if r[0] == c))
        for c in CATE_ORDER if any(r[0] == c for r in rows)))
    ws['A2'].font = Font(size=10, color='666666', name='微软雅黑')

    ws.append([])
    heads = ['序号', '类别', '名称', '相对路径', '大小', '说明']
    r = 4
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=r, column=j, value=h)
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, CENTER, BORDER
    ws.row_dimensions[r].height = 22
    r += 1
    seq = 0
    last_cate = None
    for cate, name, rel, size, note in rows:
        seq += 1
        show = cate if cate != last_cate else ''
        last_cate = cate
        for j, v in enumerate([seq, show, name, rel, size, note], start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = LEFT if j in (3, 4, 6) else CENTER
        r += 1
    autofit(ws, len(heads), [6, 14, 52, 46, 10, 40])
    ws.freeze_panes = 'A5'
    fp = os.path.join(OUT, '附件4_交付物清单.xlsx')
    wb.save(fp)
    return fp, len(rows)


def main():
    os.makedirs(OUT, exist_ok=True)
    # 清理已取消的旧附件（v2.17.0 起不再产出）
    for old in ('附件3_热线与纳统小区全量匹配表.xlsx',
                '附件4_纳统匹配候选与未匹配清单.xlsx'):
        p = os.path.join(OUT, old)
        if os.path.exists(p):
            os.remove(p)
            print('[x] 已移除旧附件 %s' % old)
    for fn, name in [(deliver_risk, '附件1 高风险小区评分全量清单'),
                     (deliver_dedup, '附件2 投诉去重建议清单'),
                     (deliver_wordcloud, '附件3 词云清单'),
                     (deliver_manifest, '附件4 交付物清单')]:
        fp, n = fn()
        print('[v] %-28s %5d 行  %s' % (name, n, os.path.relpath(fp, ROOT)))
    return 0


if __name__ == '__main__':
    sys.exit(main())
